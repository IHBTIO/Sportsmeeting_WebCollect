# -*- coding: utf-8 -*-
"""app.py v4.9 - 增加端口自动避让、改密接口"""
import io, os, sys, socket
from datetime import datetime
from functools import wraps
from flask import Flask, request, jsonify, render_template, send_file, session

import database as db

def resource_path(rel):
    base = getattr(sys, "_MEIPASS", None) or os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, rel)

app = Flask(__name__, template_folder=resource_path("templates"), static_folder=resource_path("static"))
app.config["JSON_AS_ASCII"] = False
try: app.json.ensure_ascii = False
except Exception: pass
app.secret_key = "sports-meet-secret-key-2025"

db.init_db()

def ok(msg="ok", **kw):
    d = {"ok": True, "message": msg}; d.update(kw); return jsonify(d), 200

def fail(msg, code=400, **kw):
    d = {"ok": False, "message": msg}; d.update(kw); return jsonify(d), code

def get_available_port(start_port=5000, max_attempts=20):
    """自动寻找可用端口"""
    for port in range(start_port, start_port + max_attempts):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(("0.0.0.0", port))
                return port
            except OSError:
                continue
    return start_port # 兜底

def lan_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try: s.connect(("8.8.8.8", 80)); return s.getsockname()[0]
    except: return "127.0.0.1"
    finally: s.close()

def admin_required(fn):
    @wraps(fn)
    def w(*a, **k):
        if not session.get("is_admin"): return fail("未登录", code=401)
        return fn(*a, **k)
    return w

@app.route("/")
def index(): return render_template("index.html")

@app.route("/query")
def query_page(): return render_template("query.html")

@app.route("/admin")
def admin_page(): return render_template("admin.html")

@app.route("/print/<class_name>")
@admin_required
def print_page(class_name):
    return render_template("class_print.html", class_name=class_name)

@app.route("/api/meta")
def api_meta():
    classes_map = {d: db.get_classes_by_department(d) for d in db.get_departments()}
    return ok(
        departments=db.get_departments(),
        classes_by_department=classes_map,
        projects=db.get_projects(),
        max_events=int(db.get_setting("max_events", "2")),
        registration_open=db.is_registration_open(),
        registration_deadline=db.get_setting("registration_deadline", ""),
        registration_notice=db.get_setting("registration_notice", "")
    )

@app.route("/api/submit", methods=["POST"])
def api_submit():
    if not db.is_registration_open(): return fail("报名已截止", code=403)
    data = request.get_json(silent=True) or {}
    dept = str(data.get("department", "")).strip()
    cls = str(data.get("class_name", "")).strip()
    if not dept or not db.is_valid_class(dept, cls): return fail("班级与学部不匹配")
    name = str(data.get("name", "")).strip()
    if not name: return fail("请填写姓名")
    gender = str(data.get("gender", "")).strip()
    raw = data.get("events", [])
    if not isinstance(raw, list) or not raw: return fail("请选择项目")
    
    events = list(dict.fromkeys([str(x).strip() for x in raw if str(x).strip()]))
    valid = db.get_projects()
    bad = [e for e in events if e not in valid]
    if bad: return fail(f"无效项目：{bad}")

    ip = request.headers.get("X-Forwarded-For", request.remote_addr or "")
    okf, msg, rid = db.add_registration(dept, cls, name, gender, events, ip)
    if not okf: return fail(msg, code=409)
    return ok(msg, id=rid, record={"class_name": cls, "name": name, "events": "、".join(events)})

@app.route("/api/revoke", methods=["POST"])
def api_revoke():
    data = request.get_json(silent=True) or {}
    try: rid = int(data.get("id"))
    except: return fail("参数错误")
    okf, msg = db.revoke_registration(rid, str(data.get("class_name","")), str(data.get("name","")))
    return ok(msg) if okf else fail(msg)

@app.route("/api/query")
def api_query():
    cls, name = request.args.get("class_name", ""), request.args.get("name", "")
    if not cls or not name: return fail("请填写完整")
    rec = db.find_registration(cls, name)
    if not rec: return fail("未找到记录", code=404)
    return ok(record={"class_name": rec["class_name"], "name": rec["name"], "gender": rec["gender"], "events": rec["events"].replace(",", "、")})

@app.route("/api/admin/login", methods=["POST"])
def api_login():
    data = request.get_json(silent=True) or {}
    real_pwd = db.get_setting("admin_password", "admin123")
    if str(data.get("password", "")) != real_pwd: return fail("密码错误", code=401)
    session["is_admin"] = True
    db.log_action("管理员登录")
    # 如果密码是默认的，告诉前端去强制修改
    need_change = (real_pwd == "admin123")
    return ok("登录成功", need_change_pwd=need_change)

@app.route("/api/admin/change_pwd", methods=["POST"])
@admin_required
def api_change_pwd():
    data = request.get_json(silent=True) or {}
    new_pwd = str(data.get("new_password", "")).strip()
    if len(new_pwd) < 6: return fail("密码至少6位")
    db.set_setting("admin_password", new_pwd)
    db.log_action("修改密码")
    return ok("密码已修改")

@app.route("/api/admin/logout", methods=["POST"])
def api_logout():
    session.pop("is_admin", None)
    return ok("已退出")

@app.route("/api/admin/session")
def api_session():
    return ok(logged_in=bool(session.get("is_admin")))

@app.route("/api/admin/data", methods=["GET"])
@admin_required
def api_admin_data():
    f = {k: request.args.get(k) for k in ("department","class_name","gender","project","keyword") if request.args.get(k)}
    recs = db.get_all_registrations(f or None)
    return ok(
        total=len(recs), records=recs, stats=db.get_statistics(f or None),
        registration_open=db.is_registration_open(), max_events=int(db.get_setting("max_events", "2")),
        notice=db.get_setting("registration_notice", ""), projects=db.get_projects()
    )

@app.route("/api/admin/settings", methods=["POST"])
@admin_required
def api_settings():
    data = request.get_json(silent=True) or {}
    if "max_events" in data: db.set_setting("max_events", data["max_events"])
    if "registration_notice" in data: db.set_setting("registration_notice", data["registration_notice"])
    if "registration_start" in data: db.set_setting("registration_start", data["registration_start"])
    if "registration_deadline" in data: db.set_setting("registration_deadline", data["registration_deadline"])
    if "registration_open" in data: db.set_setting("registration_open", "1" if data["registration_open"] else "0")
    db.log_action("修改设置")
    return ok("已保存")

@app.route("/api/admin/projects", methods=["POST", "DELETE"])
@admin_required
def api_projects():
    data = request.get_json(silent=True) or {}
    if request.method == "POST": okf, msg = db.add_project(data.get("name"))
    else: okf, msg = db.delete_project(data.get("name"))
    return (ok if okf else fail)(msg, projects=db.get_projects())

@app.route("/api/admin/delete_record", methods=["POST"])
@admin_required
def api_delete_record():
    data = request.get_json(silent=True) or {}
    db.delete_record(data.get("id"))
    db.log_action("删除记录")
    return ok("已删除")

@app.route("/api/admin/clear_data", methods=["POST"])
@admin_required
def api_clear_data():
    db.clear_all_records()
    db.log_action("一键清空")
    return ok("已清空")

@app.route("/api/admin/logs")
@admin_required
def api_logs(): return ok(logs=db.get_logs())

@app.route("/api/export")
@admin_required
def api_export():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except: return fail("缺少 openpyxl", code=500)
    f = {k: request.args.get(k) for k in ("department","class_name","gender","project","keyword") if request.args.get(k)}
    recs = db.get_all_registrations(f or None)
    
    wb = Workbook(); ws = wb.active; ws.title = "报名记录"
    ws.append(["序号", "学部", "班级", "姓名", "性别", "参赛项目", "项目数", "提交时间"])
    for c in range(1, 9): ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF"); ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="4472C4")
    for i, r in enumerate(recs, 1):
        ws.append([i, r["department"], r["class_name"], r["name"], r["gender"], r["events"].replace(",", "、"), r["event_count"], r["created_at"]])
    
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"报名表_{datetime.now().strftime('%Y%m%d')}.xlsx")

if __name__ == "__main__":
    port = get_available_port(5000)
    ip = lan_ip()
    print("="*50)
    print(f"系统已启动！")
    print(f"局域网访问地址: http://{ip}:{port}")
    print(f"后台管理: http://{ip}:{port}/admin")
    print("="*50)
    app.run(host="0.0.0.0", port=port)